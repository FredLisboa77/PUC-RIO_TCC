"""A planilha de anotação do gabarito (`.xlsx`) e o CSV oficial.

A planilha é cega: lista todos os candidatos com os mesmos fatos neutros, sem
dizer quais a ferramenta marcou. Fica protegida, com só as colunas de anotação
livres, para que um fato ou um ID não seja editado por engano.
"""

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel

import core.rules.todas  # noqa: F401 — registra as regras
from core.rules.registry import REGISTRO
from gabarito.candidatos import COLUNAS_DE_FATOS, Candidato, ErroDeGabarito

VEREDITOS = ("viola", "não viola", "indeterminado")
COLUNAS_EDITAVEIS = ("veredito", "justificativa", "referencia")
LINHA_CABECALHO = 4
LARGURAS = {"id": 45, "projeto": 9, "objeto": 40, "expressao": 80, "divisoes": 60, "veredito": 14, "justificativa": 50, "referencia": 30}
"""Largura de cada coluna; os fatos que não estão aqui ficam com `LARGURA_PADRAO`."""
LARGURA_PADRAO = 20

PERGUNTAS = {
    "MOD-001": "Esta tabela foi gerada pelo Tempo automático de data/hora?",
    "MOD-002": "Este relacionamento filtra nos dois sentidos sem que o cenário exija isso, segundo a orientação de relacionamentos bidirecionais?",
    "MOD-003": "Esta tabela deveria se relacionar com o resto do modelo, segundo o star schema? (Tabela de medidas ou de parâmetros isolada de propósito: não viola.)",
    "MOD-005": "Esta tabela funciona como dimensão de data e não está marcada como tabela de data?",
    "MOD-006": "Esta tabela é o elo intermediário de uma dimensão em floco de neve que deveria ser desnormalizada?",
    "MOD-007": "Este relacionamento é um-para-um entre tabelas que descrevem a mesma entidade?",
    "PERF-001": "Esta coluna calculada poderia ser criada no Power Query ou na origem, como a orientação de redução de dados recomenda?",
    "PERF-003": "Esta coluna de ponto flutuante é somada de forma que a imprecisão do tipo possa afetar o resultado?",
    "DAX-001": "Alguma divisão desta expressão tem denominador que pode ser zero ou BLANK, caso em que a documentação recomenda DIVIDE?",
}

INSTRUCOES = [
    "Gabarito das regras em P1–P7 — protocolo de anotação",
    "",
    "1. Julgue cada linha pela pergunta no topo da aba e pela página do Learn indicada, não por preferência pessoal.",
    "2. veredito: 'viola', 'não viola' ou 'indeterminado' (use a lista suspensa).",
    "3. 'indeterminado' só quando o modelo não traz a informação necessária (ex.: a distribuição dos valores de uma coluna).",
    "4. Justificativa obrigatória em 'viola' e 'indeterminado'.",
    "5. Referência: a URL de qualquer fonte usada além da página indicada no topo da aba.",
    "6. A planilha não mostra o que a ferramenta decidiu, de propósito: o julgamento é cego.",
    "7. Para conferir o andamento: python -m gabarito importar --parcial",
]


def _aba_de_regra(livro: Workbook, id_regra: str, candidatos: list[Candidato]) -> None:
    aba = livro.create_sheet(id_regra)
    aba["A1"] = PERGUNTAS[id_regra]
    aba["A1"].font = Font(bold=True)
    aba["A2"] = REGISTRO.meta(id_regra).url_canonica
    cabecalho = ["id", "projeto", "objeto", *COLUNAS_DE_FATOS[id_regra], *COLUNAS_EDITAVEIS]
    for coluna, nome in enumerate(cabecalho, start=1):
        celula = aba.cell(row=LINHA_CABECALHO, column=coluna, value=nome)
        celula.font = Font(bold=True)

    primeira_editavel = len(cabecalho) - len(COLUNAS_EDITAVEIS) + 1
    linhas = sorted(candidatos, key=lambda c: (c.projeto, c.objeto))
    for i, c in enumerate(linhas, start=LINHA_CABECALHO + 1):
        valores = [c.id, c.projeto, c.objeto, *(c.fatos[f] for f in COLUNAS_DE_FATOS[id_regra])]
        for coluna, valor in enumerate(valores, start=1):
            celula = aba.cell(row=i, column=coluna, value=ILLEGAL_CHARACTERS_RE.sub("", valor))
            # Texto, sempre: uma expressão DAX que começa com `=` viraria fórmula do Excel.
            celula.data_type = "s"
            celula.alignment = Alignment(wrap_text=True, vertical="top")
        for coluna in range(primeira_editavel, len(cabecalho) + 1):
            aba.cell(row=i, column=coluna).protection = Protection(locked=False)

    if linhas:
        letra = aba.cell(row=LINHA_CABECALHO, column=primeira_editavel).column_letter
        validacao = DataValidation(type="list", formula1=f'"{",".join(VEREDITOS)}"', allow_blank=True)
        validacao.add(f"{letra}{LINHA_CABECALHO + 1}:{letra}{LINHA_CABECALHO + len(linhas)}")
        aba.add_data_validation(validacao)
    for coluna, nome in enumerate(cabecalho, start=1):
        aba.column_dimensions[aba.cell(row=LINHA_CABECALHO, column=coluna).column_letter].width = LARGURAS.get(nome, LARGURA_PADRAO)
    ultima = aba.cell(row=LINHA_CABECALHO, column=len(cabecalho)).column_letter
    aba.auto_filter.ref = f"A{LINHA_CABECALHO}:{ultima}{LINHA_CABECALHO + max(len(linhas), 1)}"
    aba.freeze_panes = aba.cell(row=LINHA_CABECALHO + 1, column=4)
    # Protegida contra editar fato ou ID, mas livre para alargar coluna, ajustar
    # linha e filtrar: senão a única saída para ler uma expressão é desproteger.
    aba.protection.formatColumns = False
    aba.protection.formatRows = False
    aba.protection.autoFilter = False
    aba.protection.sheet = True


def gerar_planilha(candidatos: list[Candidato], destino: Path) -> None:
    if destino.exists():
        raise ErroDeGabarito(f"{destino} já existe e pode ter vereditos — apague ou renomeie antes de gerar outra")
    ids = [c.id for c in candidatos]
    duplicados = sorted({i for i in ids if ids.count(i) > 1})
    if duplicados:
        raise ErroDeGabarito(f"id duplicado entre os candidatos — o cruzamento contaria em dobro: {', '.join(duplicados)}")
    livro = Workbook()
    instrucoes = livro.active
    instrucoes.title = "Instruções"
    for i, linha in enumerate(INSTRUCOES, start=1):
        instrucoes.cell(row=i, column=1, value=linha)
    for id_regra in sorted(COLUNAS_DE_FATOS):
        _aba_de_regra(livro, id_regra, [c for c in candidatos if c.id_regra == id_regra])
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)


CAMPOS_CSV = ("id", "projeto", "id_regra", "objeto", "veredito", "justificativa", "referencia")


class LinhaGabarito(BaseModel):
    id: str
    projeto: str
    id_regra: str
    objeto: str
    veredito: str | None = None
    justificativa: str | None = None
    referencia: str | None = None


class Problema(BaseModel):
    aba: str
    linha: int | None
    """Número da linha na aba; `None` quando o candidato não está na planilha."""
    id: str
    motivo: str


def _texto(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor)
    return texto if texto.strip() else None


def ler_planilha(caminho: Path) -> list[tuple[str, int, LinhaGabarito]]:
    livro = load_workbook(caminho, read_only=True)
    lidas: list[tuple[str, int, LinhaGabarito]] = []
    # Pelas abas de regra, pelo nome: uma aba de notas acrescentada ou a ordem
    # trocada não mudam nada; uma aba de regra renomeada ou apagada é erro.
    faltando = [id_regra for id_regra in sorted(COLUNAS_DE_FATOS) if id_regra not in livro.sheetnames]
    if faltando:
        livro.close()
        raise ErroDeGabarito(f"aba(s) de regra ausente(s) na planilha: {', '.join(faltando)} — foram renomeadas ou apagadas?")
    for aba in (livro[id_regra] for id_regra in sorted(COLUNAS_DE_FATOS)):
        linhas = list(aba.iter_rows(min_row=LINHA_CABECALHO, values_only=True))
        if not linhas:
            continue
        cabecalho = list(linhas[0])
        sem_coluna = [nome for nome in ("id", "projeto", "objeto", *COLUNAS_EDITAVEIS) if nome not in cabecalho]
        if sem_coluna:
            livro.close()
            raise ErroDeGabarito(f"aba {aba.title}: coluna(s) ausente(s) no cabeçalho da linha {LINHA_CABECALHO}: {', '.join(sem_coluna)}")
        posicao = {nome: cabecalho.index(nome) for nome in ("id", "projeto", "objeto", *COLUNAS_EDITAVEIS)}
        for numero, valores in enumerate(linhas[1:], start=LINHA_CABECALHO + 1):
            if valores[posicao["id"]] is None:
                continue
            lidas.append(
                (
                    aba.title,
                    numero,
                    LinhaGabarito(
                        id=str(valores[posicao["id"]]),
                        projeto=str(valores[posicao["projeto"]]),
                        id_regra=aba.title,
                        objeto=str(valores[posicao["objeto"]]),
                        veredito=_texto(valores[posicao["veredito"]]),
                        justificativa=_texto(valores[posicao["justificativa"]]),
                        referencia=_texto(valores[posicao["referencia"]]),
                    ),
                )
            )
    livro.close()
    return lidas


def validar(lidas: list[tuple[str, int, LinhaGabarito]], candidatos: list[Candidato], parcial: bool) -> list[Problema]:
    """Não normaliza: um veredito fora da lista é apontado, não corrigido em silêncio."""
    problemas: list[Problema] = []
    for aba, numero, linha in lidas:
        if linha.veredito is None:
            if not parcial:
                problemas.append(Problema(aba=aba, linha=numero, id=linha.id, motivo="sem veredito"))
            continue
        if linha.veredito not in VEREDITOS:
            problemas.append(
                Problema(aba=aba, linha=numero, id=linha.id, motivo=f"veredito fora da lista: {linha.veredito!r}")
            )
        elif linha.veredito != "não viola" and linha.justificativa is None:
            problemas.append(Problema(aba=aba, linha=numero, id=linha.id, motivo="sem justificativa"))

    vistos: set[str] = set()
    for aba, numero, linha in lidas:
        if linha.id in vistos:
            problemas.append(Problema(aba=aba, linha=numero, id=linha.id, motivo="id duplicado"))
        vistos.add(linha.id)

    na_planilha = {linha.id: (aba, numero) for aba, numero, linha in lidas}
    atuais = {c.id: c for c in candidatos}
    for id_, (aba, numero) in na_planilha.items():
        if id_ not in atuais:
            problemas.append(Problema(aba=aba, linha=numero, id=id_, motivo="não é mais candidato"))
    for id_, c in atuais.items():
        if id_ not in na_planilha:
            problemas.append(Problema(aba=c.id_regra, linha=None, id=id_, motivo="candidato ausente da planilha"))
    return problemas


def progresso(lidas: list[tuple[str, int, LinhaGabarito]]) -> dict[str, tuple[int, int]]:
    contagem: dict[str, tuple[int, int]] = {}
    for aba, _, linha in lidas:
        julgadas, total = contagem.get(aba, (0, 0))
        contagem[aba] = (julgadas + (linha.veredito is not None), total + 1)
    return contagem


def gravar_csv(linhas: list[LinhaGabarito], destino: Path) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    with destino.open("w", encoding="utf-8", newline="") as arquivo:
        escritor = csv.DictWriter(arquivo, fieldnames=CAMPOS_CSV, lineterminator="\n")
        escritor.writeheader()
        for linha in sorted(linhas, key=lambda l: l.id):
            escritor.writerow({campo: getattr(linha, campo) or "" for campo in CAMPOS_CSV})


def ler_csv(caminho: Path) -> list[LinhaGabarito]:
    with caminho.open(encoding="utf-8", newline="") as arquivo:
        return [
            LinhaGabarito(**{campo: (valor or None) for campo, valor in linha.items()})
            for linha in csv.DictReader(arquivo)
        ]
