"""A planilha cega da rubrica, a importação e o CSV oficial.

Mostra o que o LLM escreveu e o que ele citou, e a referência da regra; esconde
os sinais automáticos (citação da âncora, divergências, cache, tempos), para que a
comparação humano × automático meça alguma coisa.
"""

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel

from avaliacao_geracao import ErroDeAvaliacao
from avaliacao_geracao.rodada import ItemDaRodada, Rodada, impressao
from core.rules.registry import Registro
from core.rules.todas import REGISTRO
from rag.avaliacao import RAIZ_REPO

VEREDITOS = ("sim", "não", "indeterminado")
ITENS = ("pertinente", "fiel", "recomendacao_correta")
LIMITE_TRECHO = 3000
PLANILHA = RAIZ_REPO / "eval" / "gabarito" / "rubrica-geracao-P1-P7.xlsx"
CSV_OFICIAL = RAIZ_REPO / "eval" / "rubrica_geracao.csv"

COLUNAS_LEITURA = (
    "id", "impressao", "projeto", "regra", "titulo_regra", "objeto", "mensagem", "trecho_do_modelo",
    "explicacao", "recomendacao", "citacoes", "referencia_regra",
)
COLUNAS_ANOTACAO = (*ITENS, "justificativa", "observacao")
LARGURAS = {"id": 30, "impressao": 10, "projeto": 12, "regra": 9, "titulo_regra": 25, "objeto": 30,
            "mensagem": 40, "trecho_do_modelo": 50, "explicacao": 60, "recomendacao": 50, "citacoes": 80, "referencia_regra": 60,
            "pertinente": 13, "fiel": 13, "recomendacao_correta": 13, "justificativa": 40, "observacao": 30}

INSTRUCOES = [
    "Rubrica das explicações do LLM em P1–P7 — protocolo de anotação",
    "",
    "Uma linha por achado com explicação gerada. Julgue o que o LLM escreveu contra os trechos citados e a",
    "referência da regra (coluna referencia_regra).",
    "",
    "pertinente — pelo menos um trecho citado é de fato a passagem da documentação que sustenta o achado",
    "  (não basta existir entre os trechos fornecidos).",
    "fiel — o que a explicação afirma é sustentado pelos trechos citados ou pelos fatos do achado (mensagem e",
    "  trecho_do_modelo, a expressão que o modelo recebeu), sem nada",
    "  inventado (números, nomes, comportamentos).",
    "recomendacao_correta — a recomendação resolve o problema neste modelo, sem erro técnico.",
    "",
    "Vereditos: 'sim', 'não' ou 'indeterminado' (lista suspensa). Justificativa obrigatória quando algum item é 'não'.",
    "Trechos longos estão cortados em 3.000 caracteres ([…]); o link leva ao texto completo.",
    "A planilha não mostra o que a ferramenta concluiu sobre a própria citação, de propósito: o julgamento é cego.",
    "Para conferir o andamento: python -m avaliacao_geracao importar --parcial",
]


class LinhaRubrica(BaseModel):
    id: str
    impressao: str
    pertinente: str | None = None
    fiel: str | None = None
    recomendacao_correta: str | None = None
    justificativa: str | None = None
    observacao: str | None = None


class Problema(BaseModel):
    linha: int | None
    id: str
    motivo: str


def _cortar(texto: str) -> str:
    return texto if len(texto) <= LIMITE_TRECHO else texto[:LIMITE_TRECHO] + " […]"


def _citacoes(r: ItemDaRodada) -> str:
    por_numero = {t.numero: t for t in r.item.contexto}
    partes = []
    for n in r.item.geracao.citacoes:
        t = por_numero[n]
        partes.append(f"[{n}] {t.titulo} — {t.caminho}\n{t.url}#{t.ancora}\n{_cortar(t.texto)}")
    return "\n\n".join(partes)


def _referencia(r: ItemDaRodada, registro: Registro) -> str:
    url = registro.meta(r.item.achado.id_regra).url_canonica
    ancora = next((t for t in r.item.contexto if t.origem == "ancora"), None)
    texto = _cortar(ancora.texto) if ancora is not None else "(trecho âncora ausente no contexto deste achado)"
    return f"{url}\n\n{texto}"


def gerar_planilha(rodada: Rodada, destino: Path, *, forcar: bool = False, registro: Registro | None = None) -> int:
    itens = sorted(
        ((chave, r) for chave, r in rodada.items() if r.item.situacao == "gerado"),
        key=lambda par: (par[1].projeto, par[1].item.achado.id_regra, par[1].item.achado.evidencia.objeto),
    )
    return escrever_planilha(itens, destino, forcar=forcar, registro=registro)


def escrever_planilha(itens: list[tuple[str, ItemDaRodada]], destino: Path, *, forcar: bool = False,
                      registro: Registro | None = None, instrucoes_extra: list[str] | None = None) -> int:
    """Uma linha por item, na ordem recebida; a primeira coluna é a chave dada."""
    if destino.exists() and not forcar:
        raise ErroDeAvaliacao(f"{destino} já existe e pode ter vereditos — apague, renomeie ou use --forcar")
    registro = REGISTRO if registro is None else registro
    livro = Workbook()
    instrucoes = livro.active
    instrucoes.title = "Instruções"
    for i, linha in enumerate([*INSTRUCOES, *(instrucoes_extra or [])], start=1):
        instrucoes.cell(row=i, column=1, value=linha)
    aba = livro.create_sheet("Rubrica")
    cabecalho = [*COLUNAS_LEITURA, *COLUNAS_ANOTACAO]
    for coluna, nome in enumerate(cabecalho, start=1):
        aba.cell(row=1, column=coluna, value=nome).font = Font(bold=True)
        aba.column_dimensions[aba.cell(row=1, column=coluna).column_letter].width = LARGURAS[nome]
    for numero, (chave, r) in enumerate(itens, start=2):
        meta = registro.meta(r.item.achado.id_regra)
        valores = [
            chave, impressao(r.item), r.projeto, meta.id, meta.titulo, r.item.achado.evidencia.objeto,
            r.item.achado.mensagem, _cortar(r.item.achado.evidencia.trecho or ""), r.item.geracao.explicacao, r.item.geracao.recomendacao,
            _citacoes(r), _referencia(r, registro),
        ]
        for coluna, valor in enumerate(valores, start=1):
            celula = aba.cell(row=numero, column=coluna, value=ILLEGAL_CHARACTERS_RE.sub("", valor))
            celula.data_type = "s"  # Review Focus 1: texto começando com `=` não vira fórmula
            celula.alignment = Alignment(wrap_text=True, vertical="top")
        for coluna in range(len(COLUNAS_LEITURA) + 1, len(cabecalho) + 1):
            celula = aba.cell(row=numero, column=coluna)
            celula.protection = Protection(locked=False)
            celula.alignment = Alignment(wrap_text=True, vertical="top")
    if itens:
        validacao = DataValidation(type="list", formula1=f'"{",".join(VEREDITOS)}"', allow_blank=True)
        for nome in ITENS:
            letra = aba.cell(row=1, column=cabecalho.index(nome) + 1).column_letter
            validacao.add(f"{letra}2:{letra}{len(itens) + 1}")
        aba.add_data_validation(validacao)
    aba.freeze_panes = "B2"
    aba.auto_filter.ref = f"A1:{aba.cell(row=1, column=len(cabecalho)).column_letter}{max(len(itens), 1) + 1}"
    aba.protection.formatColumns = False
    aba.protection.formatRows = False
    aba.protection.autoFilter = False
    aba.protection.sheet = True
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return len(itens)


def _texto(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def ler_planilha(caminho: Path) -> list[tuple[int, LinhaRubrica]]:
    livro = load_workbook(caminho, read_only=True)
    if "Rubrica" not in livro.sheetnames:
        livro.close()
        raise ErroDeAvaliacao("aba 'Rubrica' ausente na planilha — foi renomeada ou apagada?")
    linhas = list(livro["Rubrica"].iter_rows(values_only=True))
    livro.close()
    if not linhas:
        return []
    cabecalho = list(linhas[0])
    faltando = [n for n in ("id", "impressao", *COLUNAS_ANOTACAO) if n not in cabecalho]
    if faltando:
        raise ErroDeAvaliacao(f"coluna(s) ausente(s) no cabeçalho: {', '.join(faltando)}")
    pos = {n: cabecalho.index(n) for n in ("id", "impressao", *COLUNAS_ANOTACAO)}
    lidas = []
    for numero, valores in enumerate(linhas[1:], start=2):
        if valores[pos["id"]] is None:
            continue
        lidas.append((numero, LinhaRubrica(
            id=str(valores[pos["id"]]), impressao=str(valores[pos["impressao"]]),
            **{n: _texto(valores[pos[n]]) for n in COLUNAS_ANOTACAO},
        )))
    return lidas


def validar(lidas: list[tuple[int, LinhaRubrica]], rodada: Rodada, parcial: bool) -> list[Problema]:
    """Não normaliza: um veredito fora da lista é apontado, não corrigido em silêncio."""
    problemas: list[Problema] = []
    vistos: set[str] = set()
    for numero, linha in lidas:
        if linha.id in vistos:
            problemas.append(Problema(linha=numero, id=linha.id, motivo="id duplicado"))
        vistos.add(linha.id)
        if linha.id not in rodada:
            problemas.append(Problema(linha=numero, id=linha.id, motivo="id não existe na rodada em disco"))
        elif impressao(rodada[linha.id].item) != linha.impressao:
            problemas.append(Problema(
                linha=numero, id=linha.id,
                motivo="impressão digital divergente: o texto julgado não é o da rodada em disco",
            ))
        vereditos = [getattr(linha, n) for n in ITENS]
        if all(v is None for v in vereditos) and parcial:
            continue
        if any(v is None for v in vereditos):
            problemas.append(Problema(linha=numero, id=linha.id, motivo="sem veredito em algum item"))
        fora = [v for v in vereditos if v is not None and v not in VEREDITOS]
        if fora:
            problemas.append(Problema(linha=numero, id=linha.id, motivo=f"veredito fora da lista: {fora!r}"))
        if "não" in vereditos and linha.justificativa is None:
            problemas.append(Problema(linha=numero, id=linha.id, motivo="'não' sem justificativa"))
    if not parcial:
        for chave in sorted({k for k, r in rodada.items() if r.item.situacao == "gerado"} - vistos):
            problemas.append(Problema(linha=None, id=chave, motivo="achado gerado da rodada ausente da planilha"))
    return problemas


CAMPOS_CSV = ("id", "impressao", *COLUNAS_ANOTACAO)


def gravar_csv(linhas: list[LinhaRubrica], caminho: Path) -> None:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", encoding="utf-8", newline="") as arquivo:
        escritor = csv.DictWriter(arquivo, fieldnames=CAMPOS_CSV, lineterminator="\n")
        escritor.writeheader()
        for linha in sorted(linhas, key=lambda l: l.id):
            escritor.writerow({c: getattr(linha, c) or "" for c in CAMPOS_CSV})


def ler_csv(caminho: Path) -> list[LinhaRubrica]:
    if not caminho.is_file():
        raise ErroDeAvaliacao(f"{caminho} não existe — rode `python -m avaliacao_geracao importar`")
    with caminho.open(encoding="utf-8", newline="") as arquivo:
        return [
            LinhaRubrica(**{c: (v or None) for c, v in linha.items()} | {"id": linha["id"], "impressao": linha["impressao"]})
            for linha in csv.DictReader(arquivo)
        ]
