"""Planilha de anotação: cega, protegida, com lista suspensa."""

import pytest
from openpyxl import load_workbook

from gabarito.candidatos import COLUNAS_DE_FATOS, Candidato, ErroDeGabarito
from gabarito.planilha import (
    LINHA_CABECALHO,
    PERGUNTAS,
    gerar_planilha,
    gravar_csv,
    ler_csv,
    ler_planilha,
    progresso,
    validar,
)


def _candidatos():
    return [
        Candidato(id_regra="MOD-002", projeto="P1", objeto="Vendas[Produto] → Produto[Produto]",
                  fatos={"direcao": "ambos_sentidos", "cardinalidades": "muitos:um", "ativo": "sim"}),
        Candidato(id_regra="PERF-001", projeto="P3", objeto="Pessoa[Salário]",
                  fatos={"expressao": "Pessoa[Base] * 1.1"}),
    ]


def _ids_da_aba(aba):
    return [aba.cell(row=r, column=1).value for r in range(LINHA_CABECALHO + 1, aba.max_row + 1)]


def test_uma_aba_por_regra_mais_instrucoes(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    livro = load_workbook(destino)
    assert livro.sheetnames == ["Instruções", *sorted(COLUNAS_DE_FATOS)]


def test_pergunta_no_topo_e_cabecalho_com_os_fatos(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    aba = load_workbook(destino)["MOD-002"]
    assert aba["A1"].value == PERGUNTAS["MOD-002"]
    assert aba["A2"].value.startswith("https://learn.microsoft.com/")
    cabecalho = [c.value for c in aba[LINHA_CABECALHO]]
    assert cabecalho == ["id", "projeto", "objeto", "direcao", "cardinalidades", "ativo", "veredito", "justificativa", "referencia"]


def test_id_com_caracteres_especiais_volta_identico(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    livro = load_workbook(destino)
    assert _ids_da_aba(livro["MOD-002"]) == ["MOD-002|P1|Vendas[Produto] → Produto[Produto]"]
    assert _ids_da_aba(livro["PERF-001"]) == ["PERF-001|P3|Pessoa[Salário]"]


def test_veredito_tem_lista_suspensa(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    aba = load_workbook(destino)["MOD-002"]
    validacoes = aba.data_validations.dataValidation
    assert len(validacoes) == 1
    assert validacoes[0].formula1 == '"viola,não viola,indeterminado"'
    assert "G5" in str(validacoes[0].sqref)


def test_aba_protegida_com_so_as_colunas_de_anotacao_livres(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    aba = load_workbook(destino)["MOD-002"]
    assert aba.protection.sheet is True
    assert aba["A5"].protection.locked is True
    assert [aba.cell(row=5, column=c).protection.locked for c in (7, 8, 9)] == [False, False, False]


def test_planilha_e_cega(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    valores = [
        str(celula.value).strip().lower()
        for aba in load_workbook(destino).worksheets
        for linha in aba.iter_rows()
        for celula in linha
        if celula.value is not None
    ]
    textos = " ".join(valores)
    for proibido in ("severidade", "achado", "mensagem"):
        assert proibido not in textos, proibido
    assert not {"alta", "media", "média", "baixa"} & set(valores)


def test_recusa_sobrescrever(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    with pytest.raises(ErroDeGabarito, match="já existe"):
        gerar_planilha(_candidatos(), destino)


def _preencher(destino, vereditos):
    """vereditos: {id: (veredito, justificativa)} — escreve como o Fred faria."""
    livro = load_workbook(destino)
    for aba in livro.worksheets[1:]:
        cabecalho = [c.value for c in aba[LINHA_CABECALHO]]
        col_v = cabecalho.index("veredito") + 1
        for r in range(LINHA_CABECALHO + 1, aba.max_row + 1):
            id_ = aba.cell(row=r, column=1).value
            if id_ in vereditos:
                veredito, justificativa = vereditos[id_]
                aba.cell(row=r, column=col_v, value=veredito)
                aba.cell(row=r, column=col_v + 1, value=justificativa)
    livro.save(destino)


ID_REL = "MOD-002|P1|Vendas[Produto] → Produto[Produto]"
ID_COL = "PERF-001|P3|Pessoa[Salário]"


def _planilha(tmp_path, vereditos):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    _preencher(destino, vereditos)
    return destino


def test_planilha_completa_valida_sem_problemas(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "bidirecional sem motivo"), ID_COL: ("não viola", None)})
    lidas = ler_planilha(destino)
    assert validar(lidas, _candidatos(), parcial=False) == []
    assert {l.id: l.veredito for _, _, l in lidas} == {ID_REL: "viola", ID_COL: "não viola"}


def test_veredito_ausente_e_problema(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "x")})
    problemas = validar(ler_planilha(destino), _candidatos(), parcial=False)
    assert [(p.id, p.motivo) for p in problemas] == [(ID_COL, "sem veredito")]


@pytest.mark.parametrize("valor", ["Viola", " não viola ", "nao viola", "talvez"])
def test_veredito_fora_da_lista_e_problema_sem_normalizar(tmp_path, valor):
    destino = _planilha(tmp_path, {ID_REL: (valor, "x"), ID_COL: ("não viola", None)})
    problemas = validar(ler_planilha(destino), _candidatos(), parcial=False)
    assert [(p.aba, p.id) for p in problemas] == [("MOD-002", ID_REL)]
    assert "fora da lista" in problemas[0].motivo


@pytest.mark.parametrize("veredito", ["viola", "indeterminado"])
def test_viola_e_indeterminado_exigem_justificativa(tmp_path, veredito):
    destino = _planilha(tmp_path, {ID_REL: (veredito, "  "), ID_COL: ("não viola", None)})
    problemas = validar(ler_planilha(destino), _candidatos(), parcial=False)
    assert [(p.id, p.motivo) for p in problemas] == [(ID_REL, "sem justificativa")]


def test_id_faltando_e_sobrando_sao_problemas(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "x"), ID_COL: ("não viola", None)})
    atuais = [_candidatos()[0], Candidato(id_regra="PERF-001", projeto="P3", objeto="Pessoa[Novo]", fatos={"expressao": "1"})]
    motivos = {(p.id, p.motivo) for p in validar(ler_planilha(destino), atuais, parcial=False)}
    assert motivos == {(ID_COL, "não é mais candidato"), ("PERF-001|P3|Pessoa[Novo]", "candidato ausente da planilha")}


def test_parcial_aceita_em_branco_e_informa_progresso(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "x")})
    lidas = ler_planilha(destino)
    assert validar(lidas, _candidatos(), parcial=True) == []
    assert progresso(lidas)["MOD-002"] == (1, 1)
    assert progresso(lidas)["PERF-001"] == (0, 1)


def test_csv_ida_e_volta(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "bidirecional, sem motivo"), ID_COL: ("não viola", None)})
    linhas = [l for _, _, l in ler_planilha(destino)]
    csv = tmp_path / "ground_truth.csv"
    gravar_csv(linhas, csv)
    assert csv.read_text(encoding="utf-8").splitlines()[0] == "id,projeto,id_regra,objeto,veredito,justificativa,referencia"
    assert ler_csv(csv) == sorted(linhas, key=lambda l: l.id)


def test_protecao_deixa_ajustar_colunas_linhas_e_filtrar(tmp_path):
    destino = tmp_path / "a.xlsx"
    gerar_planilha(_candidatos(), destino)
    aba = load_workbook(destino)["PERF-001"]
    assert aba.protection.sheet is True
    assert (aba.protection.formatColumns, aba.protection.formatRows, aba.protection.autoFilter) == (False, False, False)
    assert aba.auto_filter.ref == "A4:G5"
    assert aba.column_dimensions["A"].width >= 40
    assert aba.column_dimensions["D"].width >= 60


def test_texto_que_parece_formula_fica_texto_e_caractere_de_controle_sai(tmp_path):
    destino = tmp_path / "a.xlsx"
    candidatos = [
        Candidato(id_regra="PERF-001", projeto="P1", objeto="T[A]", fatos={"expressao": "=SUM(T[B])"}),
        Candidato(id_regra="PERF-001", projeto="P1", objeto="T[B]", fatos={"expressao": "1\x07 + 2"}),
    ]
    gerar_planilha(candidatos, destino)
    aba = load_workbook(destino)["PERF-001"]
    assert aba["D5"].value == "=SUM(T[B])" and aba["D5"].data_type == "s"
    assert aba["D6"].value == "1 + 2"


def test_gerar_recusa_id_duplicado(tmp_path):
    dobrado = _candidatos() + [_candidatos()[0]]
    with pytest.raises(ErroDeGabarito, match="duplicado"):
        gerar_planilha(dobrado, tmp_path / "a.xlsx")


def test_aba_extra_e_ignorada_e_aba_faltando_vira_erro(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "x"), ID_COL: ("não viola", None)})
    livro = load_workbook(destino)
    livro.create_sheet("Notas", 0)["A1"] = "rascunho do Fred"
    livro.save(destino)
    assert {l.id for _, _, l in ler_planilha(destino)} == {ID_REL, ID_COL}

    livro = load_workbook(destino)
    livro["MOD-002"].title = "MOD-002 (revisar)"
    livro.save(destino)
    with pytest.raises(ErroDeGabarito, match="MOD-002"):
        ler_planilha(destino)


def test_coluna_do_cabecalho_apagada_vira_erro(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "x")})
    livro = load_workbook(destino)
    aba = livro["MOD-002"]
    cabecalho = [c.value for c in aba[LINHA_CABECALHO]]
    aba.cell(row=LINHA_CABECALHO, column=cabecalho.index("veredito") + 1).value = None
    livro.save(destino)
    with pytest.raises(ErroDeGabarito, match="veredito"):
        ler_planilha(destino)


def test_validar_aponta_id_duplicado_e_candidato_ausente_sem_numero_de_linha(tmp_path):
    destino = _planilha(tmp_path, {ID_REL: ("viola", "x"), ID_COL: ("não viola", None)})
    lidas = ler_planilha(destino)
    lidas.append(lidas[0])
    atuais = _candidatos() + [Candidato(id_regra="PERF-001", projeto="P3", objeto="Pessoa[Novo]", fatos={"expressao": "1"})]
    problemas = validar(lidas, atuais, parcial=False)
    assert (ID_REL, "id duplicado") in {(p.id, p.motivo) for p in problemas}
    ausente = next(p for p in problemas if p.motivo == "candidato ausente da planilha")
    assert ausente.linha is None
