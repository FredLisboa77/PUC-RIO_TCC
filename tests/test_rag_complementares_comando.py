"""O comando python -m rag.complementares, com um ambiente sintético."""

import json

import pytest
from fabricas_analise import trecho_indexado
from openpyxl import load_workbook

from rag import complementares as comando
from rag.avaliacao import Consulta
from rag.pertinencia import LINHA_CABECALHO

ANCORA = trecho_indexado("learn:ancora", "auto date time")
TRECHOS = [ANCORA, trecho_indexado("learn:a", "a"), trecho_indexado("learn:b", "b")]


def _ambiente():
    return comando.Ambiente(
        consultas=[Consulta(id="REGRA/MOD-001", tipo="regra", id_regra="MOD-001", texto="q")],
        ancoras={"MOD-001": ANCORA},
        trechos=TRECHOS,
        ranqueadores={
            "bm25": lambda q, k: [0, 2][:k],
            "densa": lambda q, k: [1, 0][:k],
            "hibrida": lambda q, k: [1, 2][:k],
        },
    )


@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    monkeypatch.setattr(comando, "ARQUIVO_PLANILHA", tmp_path / "gabarito" / "p.xlsx")
    monkeypatch.setattr(comando, "ARQUIVO_CSV", tmp_path / "pertinencia.csv")
    monkeypatch.setattr(comando, "PASTA_RESULTADOS", tmp_path / "results")
    monkeypatch.setattr(comando, "carregar_ambiente", _ambiente)
    return tmp_path


def _julgar(caminho, vereditos):
    livro = load_workbook(caminho)
    aba = livro["MOD-001"]
    cab = [c.value for c in aba[LINHA_CABECALHO]]
    for r in range(LINHA_CABECALHO + 1, aba.max_row + 1):
        aba.cell(row=r, column=cab.index("veredito") + 1).value = vereditos[aba.cell(row=r, column=1).value]
    livro.save(caminho)


def test_fluxo_completo_decide_e_grava_json(ambiente, capsys):
    assert comando.main(["gerar"]) == 0
    _julgar(ambiente / "gabarito" / "p.xlsx", {"MOD-001|learn:a#s": "pertinente", "MOD-001|learn:b#s": "não pertinente"})
    assert comando.main(["importar"]) == 0
    assert comando.main(["medir"]) == 0
    dados = json.loads(next((ambiente / "results").glob("complementares-*.json")).read_text(encoding="utf-8"))
    assert dados["por_busca"]["bm25"]["macro"]["principal"] == 0.0
    assert dados["por_busca"]["densa"]["macro"]["principal"] == 1.0
    assert dados["decisao"]["vencedora"] == "densa"
    capsys.readouterr().out.encode("ascii")


def test_importar_incompleto_nao_grava(ambiente, capsys):
    comando.main(["gerar"])
    assert comando.main(["importar"]) == 1
    assert "sem veredito" in capsys.readouterr().err
    assert not (ambiente / "pertinencia.csv").exists()


def test_medir_sem_csv_indica_importar(ambiente, capsys):
    assert comando.main(["medir"]) == 1
    assert "importar" in capsys.readouterr().err


def test_sem_cache_de_embeddings_indica_a_medicao(monkeypatch, capsys):
    def sem_cache():
        raise comando.ErroDeComplementares("sem cache de embeddings para secao-sem - rode `python -m rag.medicao`")

    monkeypatch.setattr(comando, "carregar_ambiente", sem_cache)
    assert comando.main(["gerar"]) == 1
    assert "rag.medicao" in capsys.readouterr().err


from rag.indice import ARQUIVO_CONFIG  # noqa: E402

SEM_BASE = not ARQUIVO_CONFIG.is_file()


@pytest.mark.skipif(SEM_BASE, reason="índice da base RAG não construído")
def test_pool_real_cobre_as_9_regras():
    amb = comando.carregar_ambiente()
    pool = comando.formar_pool(amb.consultas, amb.ancoras, amb.trechos, amb.ranqueadores)
    assert len({p.id_regra for p in pool}) == 9
