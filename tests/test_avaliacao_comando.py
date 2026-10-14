"""CLI da avaliação: fluxo completo com LLM falso e mensagens em vez de traceback."""

import json

import pytest
from openpyxl import load_workbook

import avaliacao_geracao.__main__ as comando
from avaliacao_geracao.planilha import gravar_csv, ler_csv
from conftest import escrever_pbip, modelo_tmsl, tabela
from core import auditoria
from core.rules.registry import Registro
from fabricas_analise import LLMFalso, achado, base_padrao, gabarito, meta


@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    registro = Registro()
    registro.registrar(meta(), lambda modelo: [achado(objeto=t.nome) for t in modelo.tabelas])
    monkeypatch.setattr(auditoria, "REGISTRO", registro)
    monkeypatch.setattr(auditoria, "carregar_gabaritos", lambda: {"MOD-001": gabarito()})
    for nome in ["PASTA_RODADA", "PLANILHA", "CSV_OFICIAL", "PASTA_RESULTADOS"]:
        destino = {"PASTA_RODADA": tmp_path / "rodada", "PLANILHA": tmp_path / "r.xlsx",
                   "CSV_OFICIAL": tmp_path / "r.csv", "PASTA_RESULTADOS": tmp_path / "results"}[nome]
        monkeypatch.setattr(comando, nome, destino)
    pbip = tmp_path / "pbip"
    escrever_pbip(pbip / "P1_A", nome="P1_A", model_bim=modelo_tmsl([tabela("T1"), tabela("T2")]))
    monkeypatch.setattr(comando, "PASTA_PBIP", pbip)
    monkeypatch.setattr(comando, "criar_llm", lambda: LLMFalso())
    monkeypatch.setattr(comando, "carregar_base", base_padrao)
    monkeypatch.setattr(comando, "commit_atual", lambda: "abc1234")
    monkeypatch.setattr(comando, "REGISTRO_DA_PLANILHA", registro)
    return tmp_path


def _anotar_tudo(caminho):
    livro = load_workbook(caminho)
    aba = livro["Rubrica"]
    cabecalho = [c.value for c in next(aba.iter_rows(min_row=1, max_row=1))]
    for linha in range(2, aba.max_row + 1):
        for nome in ["pertinente", "fiel", "recomendacao_correta"]:
            aba.cell(row=linha, column=cabecalho.index(nome) + 1, value="sim")
    livro.save(caminho)


def _importar_tudo(ambiente):
    assert comando.main(["auditar"]) == 0
    assert comando.main(["gerar"]) == 0
    _anotar_tudo(ambiente / "r.xlsx")
    assert comando.main(["importar"]) == 0


def test_fluxo_completo(ambiente, capsys):
    assert comando.main(["auditar"]) == 0
    assert "fora do cache: 2" in capsys.readouterr().out
    assert comando.main(["gerar"]) == 0
    assert comando.main(["gerar"]) == 1  # não sobrescreve
    assert comando.main(["importar"]) == 1  # sem vereditos
    assert comando.main(["importar", "--parcial"]) == 0
    _anotar_tudo(ambiente / "r.xlsx")
    assert comando.main(["importar"]) == 0
    assert comando.main(["medir"]) == 0
    geracao = next((ambiente / "results").glob("geracao-*.json"))
    assert json.loads(geracao.read_text(encoding="utf-8"))["utilizavel"]["total"]["principal"] == 1.0
    capsys.readouterr().out.encode("ascii")


def test_auditar_com_ollama_fora_do_ar_sai_com_mensagem(ambiente, monkeypatch, capsys):
    monkeypatch.setattr(comando, "criar_llm", lambda: LLMFalso(falha_no_aquecimento=True))
    assert comando.main(["auditar"]) == 1
    assert "aquecimento" in capsys.readouterr().err


def test_medir_sem_csv_e_consolidar_sem_origem_saem_com_mensagem(ambiente, capsys):
    assert comando.main(["medir"]) == 1
    assert comando.main(["consolidar"]) == 1
    assert "erro:" in capsys.readouterr().err


@pytest.mark.parametrize("alteracao", [{"id": "P9_X-999"}, {"impressao": "00000000"}])
def test_medir_recusa_csv_que_nao_bate_com_a_rodada(ambiente, capsys, alteracao):
    _importar_tudo(ambiente)
    capsys.readouterr()
    linhas = ler_csv(ambiente / "r.csv")
    linhas[0] = linhas[0].model_copy(update=alteracao)
    gravar_csv(linhas, ambiente / "r.csv")
    assert comando.main(["medir"]) == 1
    err = capsys.readouterr().err
    assert "nada medido" in err
    assert not list((ambiente / "results").glob("geracao-*.json"))


def test_consolidar_com_json_sem_chave_sai_com_mensagem(ambiente, capsys):
    (ambiente / "results").mkdir()
    (ambiente / "results" / "geracao-2026-01-01.json").write_text("{}", encoding="utf-8")
    assert comando.main(["consolidar"]) == 1
    assert "erro:" in capsys.readouterr().err
