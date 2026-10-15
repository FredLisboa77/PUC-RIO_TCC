"""Teste do glossário na instrução (protocolo de 15/10/2026): análise de erros."""

import pytest
from openpyxl import load_workbook

from avaliacao_geracao.estatistica import mcnemar_exato
from avaliacao_geracao.glossario import (
    GLOSSARIO, SISTEMA_COM_GLOSSARIO, ClienteComInstrucao, gerar_planilha_cega, indicador_lexico, medir_glossario,
    montar_amostra,
)
from avaliacao_geracao.planilha import LinhaRubrica
from avaliacao_geracao.rodada import ItemDaRodada, id_do_achado
from core.analise.base import AchadoEnriquecido, Fidelidade, Geracao
from core.analise.llm import RespostaLLM
from core.analise.prompt import SISTEMA
from fabricas_analise import achado, contexto


class _Falso:
    nome = "modelo-x"

    def __init__(self):
        self.recebidas = None
        self.aquecido = False

    def aquecer(self):
        self.aquecido = True

    def gerar(self, mensagens, esquema):
        self.recebidas = mensagens
        return RespostaLLM(texto="{}", segundos=0.1)


def test_instrucao_com_glossario_estende_a_original():
    assert SISTEMA_COM_GLOSSARIO.startswith(SISTEMA)
    assert GLOSSARIO in SISTEMA_COM_GLOSSARIO
    assert 'Nunca traduza como "compreender"' in GLOSSARIO


def test_cliente_troca_so_a_mensagem_de_sistema_e_nao_altera_a_lista_recebida():
    falso = _Falso()
    cliente = ClienteComInstrucao(falso, "NOVA")
    mensagens = [{"role": "system", "content": SISTEMA}, {"role": "user", "content": "achado"}]
    cliente.aquecer()
    cliente.gerar(mensagens, {})
    assert falso.recebidas == [{"role": "system", "content": "NOVA"}, {"role": "user", "content": "achado"}]
    assert mensagens[0]["content"] == SISTEMA
    assert cliente.nome == "modelo-x" and falso.aquecido


def test_cliente_recusa_mensagens_sem_instrucao_de_sistema():
    with pytest.raises(ValueError):
        ClienteComInstrucao(_Falso(), "NOVA").gerar([{"role": "user", "content": "x"}], {})


def _item(objeto, regra="MOD-001", explicacao="Texto.", situacao="gerado"):
    a = achado(objeto=objeto).model_copy(update={"id_regra": regra})
    if situacao != "gerado":
        return AchadoEnriquecido(achado=a, contexto=contexto(2), situacao=situacao, motivo_descarte="erro_llm", modelo="m")
    return AchadoEnriquecido(
        achado=a, contexto=contexto(2), situacao="gerado", modelo="m", cita_ancora=True, fidelidade=Fidelidade(),
        geracao=Geracao(explicacao=explicacao, recomendacao="Faça.", citacoes=[1]),
    )


def _rodada(itens):
    return {id_do_achado("P1_A", i): ItemDaRodada("P1_A", i) for i in itens}


def test_amostra_e_deterministica_opaca_e_completa():
    original = _rodada([_item(f"T{i}") for i in range(10)])
    nova = _rodada([_item(f"T{i}", explicacao="Novo.") for i in range(10)])
    chave = montar_amostra(original, nova, n_originais=4, semente=7)
    assert chave == montar_amostra(original, nova, n_originais=4, semente=7)
    assert list(chave) == [f"R{n:02d}" for n in range(1, 15)]
    versoes = [v["versao"] for v in chave.values()]
    assert versoes.count("glossario") == 10 and versoes.count("original") == 4
    assert {v["achado"] for v in chave.values() if v["versao"] == "glossario"} == set(nova)
    assert chave != montar_amostra(original, nova, n_originais=4, semente=8)


def test_amostra_deixa_fora_achado_sem_explicacao_nova():
    original = _rodada([_item("T0"), _item("T1")])
    nova = _rodada([_item("T0"), _item("T1", situacao="sem_fundamentacao")])
    chave = montar_amostra(original, nova, n_originais=1, semente=1)
    assert sum(v["versao"] == "glossario" for v in chave.values()) == 1


def test_planilha_cega_mostra_codigo_e_nao_a_versao(tmp_path):
    original = _rodada([_item("T0", explicacao="Antigo.")])
    nova = _rodada([_item("T0", explicacao="Novo.")])
    chave = montar_amostra(original, nova, n_originais=1, semente=3)
    destino = tmp_path / "cega.xlsx"
    assert gerar_planilha_cega(chave, original, nova, destino) == 2
    celulas = [str(c.value) for linha in load_workbook(destino)["Rubrica"].iter_rows() for c in linha if c.value]
    texto = " ".join(celulas)
    assert "Antigo." in texto and "Novo." in texto
    assert "R01" in celulas and "R02" in celulas
    for proibido in ["glossario", "original", "P1_A|MOD-001|T0"]:
        assert proibido not in texto


def test_mcnemar_exato_bilateral():
    assert mcnemar_exato(0, 0) == 1.0
    assert mcnemar_exato(0, 6) == pytest.approx(0.03125)
    assert mcnemar_exato(1, 8) == pytest.approx(0.0390625)
    assert mcnemar_exato(5, 5) == 1.0


def _linha(id_, p="sim", f="sim", r="sim"):
    return LinhaRubrica(id=id_, impressao="x", pertinente=p, fiel=f, recomendacao_correta=r,
                        justificativa=None if "não" not in (p, f, r) else "motivo")


def test_medida_pareada_conta_descarte_como_nao_utilizavel_e_mede_consistencia():
    ids = [f"P1_A|PERF-001|C{i}" for i in range(3)] + ["P1_A|DAX-001|M"]
    originais = [_linha(ids[0], f="não"), _linha(ids[1], f="não"), _linha(ids[2]), _linha(ids[3], f="não")]
    chave = {
        "R01": {"versao": "glossario", "achado": ids[0]},
        "R02": {"versao": "glossario", "achado": ids[1]},
        "R03": {"versao": "glossario", "achado": ids[2]},
        "R04": {"versao": "original", "achado": ids[0]},
        "R05": {"versao": "original", "achado": ids[2]},
    }
    novas = [_linha("R01"), _linha("R02"), _linha("R03"), _linha("R04", f="não"), _linha("R05", f="não")]
    r = medir_glossario(originais, novas, chave, achados=ids)
    u = r["itens"]["utilizavel"]
    assert u["original"]["sim"] == 1 and u["glossario"]["sim"] == 3 and r["n_pares"] == 4
    assert u["discordantes"] == {"so_original": 0, "so_glossario": 2}
    assert u["p_mcnemar"] == pytest.approx(0.5)
    assert r["por_regra"]["DAX-001"]["utilizavel"] == {"original": "0 / 1", "glossario": "0 / 1"}
    c = r["consistencia"]["fiel"]
    assert c["n"] == 2 and c["concordancia"] == pytest.approx(0.5)
    assert r["decisao"] == {"melhora": False, "fragil": True}


def test_indicador_lexico_conta_explicacoes_da_regra_com_o_erro_de_traducao():
    rodada = _rodada([
        _item("A", regra="PERF-001", explicacao="O VertiPaq não Compreende bem a coluna."),
        _item("B", regra="PERF-001", explicacao="Gera descompressão de dados."),
        _item("C", regra="PERF-001", explicacao="Comprime pior."),
        _item("D", regra="DAX-001", explicacao="compreende"),
        _item("E", regra="PERF-001", situacao="sem_fundamentacao"),
    ])
    assert indicador_lexico(rodada) == {"com_erro": 2, "explicacoes": 3}
