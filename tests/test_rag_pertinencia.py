"""Planilha cega de pertinência dos complementares."""

import pytest
from fabricas_analise import trecho_indexado
from openpyxl import load_workbook

from rag.complementares import ErroDeComplementares, ItemPool
from rag.pertinencia import LINHA_CABECALHO, gerar_planilha, gravar_csv, ler_csv, ler_planilha, progresso, validar

ANCORAS = {"PERF-001": trecho_indexado("learn:reducao", "Prefer custom columns in Power Query.")}
POOL = [
    ItemPool(id="PERF-001|dax:rankx#s", id_regra="PERF-001", trecho=trecho_indexado("dax:rankx", "RANKX returns...")),
    ItemPool(id="PERF-001|learn:pq#s", id_regra="PERF-001", trecho=trecho_indexado("learn:pq", "=Custom column " + "z" * 900)),
]


def _gerar(tmp_path):
    destino = tmp_path / "p.xlsx"
    gerar_planilha(POOL, ANCORAS, destino)
    return destino


def _preencher(destino, vereditos):
    livro = load_workbook(destino)
    aba = livro["PERF-001"]
    cabecalho = [c.value for c in aba[LINHA_CABECALHO]]
    cv = cabecalho.index("veredito") + 1
    for r in range(LINHA_CABECALHO + 1, aba.max_row + 1):
        id_ = aba.cell(row=r, column=1).value
        if id_ in vereditos:
            aba.cell(row=r, column=cv).value = vereditos[id_][0]
            aba.cell(row=r, column=cv + 1).value = vereditos[id_][1]
    livro.save(destino)


def test_uma_aba_por_regra_com_a_ancora_no_topo_e_cega(tmp_path):
    livro = load_workbook(_gerar(tmp_path))
    assert livro.sheetnames == ["Instruções", "PERF-001"]
    aba = livro["PERF-001"]
    assert "Prefer custom columns" in aba["A3"].value
    textos = " ".join(str(c.value).lower() for linha in aba.iter_rows() for c in linha if c.value)
    for proibido in ("bm25", "densa", "hibrida", "híbrida", "combinada"):
        assert proibido not in textos


def test_texto_que_comeca_com_igual_fica_texto_e_inicio_e_cortado(tmp_path):
    aba = load_workbook(_gerar(tmp_path))["PERF-001"]
    cabecalho = [c.value for c in aba[LINHA_CABECALHO]]
    celula = aba.cell(row=LINHA_CABECALHO + 2, column=cabecalho.index("inicio_do_texto") + 1)
    assert celula.data_type == "s" and celula.value.startswith("=Custom")
    assert len(celula.value) <= 600 + len(" […]")


def test_recusa_sobrescrever(tmp_path):
    destino = _gerar(tmp_path)
    with pytest.raises(ErroDeComplementares, match="já existe"):
        gerar_planilha(POOL, ANCORAS, destino)


def test_completa_valida_e_csv_ida_e_volta(tmp_path):
    destino = _gerar(tmp_path)
    _preencher(destino, {POOL[0].id: ("não pertinente", None), POOL[1].id: ("pertinente", None)})
    lidas = ler_planilha(destino)
    assert validar(lidas, POOL, parcial=False) == []
    csv = tmp_path / "pertinencia.csv"
    gravar_csv([l for _, _, l in lidas], csv)
    assert ler_csv(csv) == {POOL[0].id: "não pertinente", POOL[1].id: "pertinente"}


@pytest.mark.parametrize("valor", ["Pertinente", " pertinente ", "talvez"])
def test_veredito_fora_da_lista_e_problema(tmp_path, valor):
    destino = _gerar(tmp_path)
    _preencher(destino, {POOL[0].id: (valor, None), POOL[1].id: ("pertinente", None)})
    assert [(p.id, "fora da lista" in p.motivo) for p in validar(ler_planilha(destino), POOL, parcial=False)] == [(POOL[0].id, True)]


def test_indeterminado_exige_justificativa_e_parcial_mostra_progresso(tmp_path):
    destino = _gerar(tmp_path)
    _preencher(destino, {POOL[0].id: ("indeterminado", None)})
    lidas = ler_planilha(destino)
    assert [(p.id, p.motivo) for p in validar(lidas, POOL, parcial=True)] == [(POOL[0].id, "sem justificativa")]
    assert progresso(lidas) == {"PERF-001": (1, 2)}


def test_id_faltando_e_sobrando(tmp_path):
    destino = _gerar(tmp_path)
    _preencher(destino, {POOL[0].id: ("pertinente", None), POOL[1].id: ("pertinente", None)})
    novo = ItemPool(id="PERF-001|learn:x#s", id_regra="PERF-001", trecho=trecho_indexado("learn:x", "x"))
    motivos = {(p.id, p.motivo) for p in validar(ler_planilha(destino), [POOL[0], novo], parcial=False)}
    assert motivos == {(POOL[1].id, "não está mais no pool"), (novo.id, "seção ausente da planilha")}
