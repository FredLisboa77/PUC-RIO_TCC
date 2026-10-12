"""O comando python -m gabarito: gerar, importar e medir, com dados falsos."""

import json

import pytest

from core.rules.base import Achado, Evidencia
from gabarito import __main__ as comando
from gabarito.candidatos import Candidato
from gabarito.planilha import LINHA_CABECALHO

CANDIDATOS = [
    Candidato(id_regra="MOD-003", projeto="P1", objeto="A", fatos={"relacionamentos": "0", "medidas": "0", "colunas": "3", "oculta": "não"}),
    Candidato(id_regra="MOD-003", projeto="P1", objeto="B", fatos={"relacionamentos": "2", "medidas": "0", "colunas": "3", "oculta": "não"}),
]


@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    monkeypatch.setattr(comando, "ARQUIVO_PLANILHA", tmp_path / "gabarito" / "anotacao.xlsx")
    monkeypatch.setattr(comando, "ARQUIVO_CSV", tmp_path / "ground_truth.csv")
    monkeypatch.setattr(comando, "PASTA_RESULTADOS", tmp_path / "results")
    monkeypatch.setattr(comando, "candidatos_do_dataset", lambda: CANDIDATOS)
    achado = Achado(id_regra="MOD-003", evidencia=Evidencia(tipo_objeto="tabela", objeto="A"), mensagem="m")
    monkeypatch.setattr(comando, "achados_dos_projetos", lambda pasta: {"P1": [achado]})
    return tmp_path


def _preencher_tudo(caminho, veredito="viola"):
    from openpyxl import load_workbook

    livro = load_workbook(caminho)
    aba = livro["MOD-003"]
    cabecalho = [c.value for c in aba[LINHA_CABECALHO]]
    col_v = cabecalho.index("veredito") + 1
    for r in range(LINHA_CABECALHO + 1, aba.max_row + 1):
        aba.cell(row=r, column=col_v, value=veredito)
        aba.cell(row=r, column=col_v + 1, value="justificativa")
    livro.save(caminho)


def test_gerar_cria_a_planilha_e_recusa_a_segunda_vez(ambiente, capsys):
    assert comando.main(["gerar"]) == 0
    assert (ambiente / "gabarito" / "anotacao.xlsx").is_file()
    assert comando.main(["gerar"]) == 1
    assert "existe e pode ter vereditos" in capsys.readouterr().err


def test_importar_incompleto_lista_problemas_e_nao_grava(ambiente, capsys):
    comando.main(["gerar"])
    assert comando.main(["importar"]) == 1
    assert "sem veredito" in capsys.readouterr().err
    assert not (ambiente / "ground_truth.csv").exists()


def test_importar_parcial_mostra_progresso(ambiente, capsys):
    comando.main(["gerar"])
    assert comando.main(["importar", "--parcial"]) == 0
    assert "MOD-003: 0/2" in capsys.readouterr().out


def test_importar_completo_grava_o_csv(ambiente):
    comando.main(["gerar"])
    _preencher_tudo(ambiente / "gabarito" / "anotacao.xlsx")
    assert comando.main(["importar"]) == 0
    assert (ambiente / "ground_truth.csv").read_text(encoding="utf-8").count("\n") == 3


def test_planilha_aberta_no_excel_vira_mensagem(ambiente, monkeypatch, capsys):
    comando.main(["gerar"])

    def bloqueada(caminho):
        raise PermissionError("arquivo em uso")

    monkeypatch.setattr(comando, "ler_planilha", bloqueada)
    assert comando.main(["importar"]) == 1
    assert "erro:" in capsys.readouterr().err


def test_medir_grava_json_e_imprime_em_ascii(ambiente, capsys):
    comando.main(["gerar"])
    _preencher_tudo(ambiente / "gabarito" / "anotacao.xlsx")
    comando.main(["importar"])
    assert comando.main(["medir"]) == 0
    resultados = list((ambiente / "results").glob("regras-*.json"))
    assert len(resultados) == 1
    dados = json.loads(resultados[0].read_text(encoding="utf-8"))
    assert dados["por_regra"]["MOD-003"]["matriz"]["vp"] == 1
    saida = capsys.readouterr().out
    saida.encode("ascii")
    assert "MOD-003" in saida and "R-12" in saida


def test_medir_sem_csv_sai_com_erro(ambiente, capsys):
    assert comando.main(["medir"]) == 1
    assert "importar" in capsys.readouterr().err
