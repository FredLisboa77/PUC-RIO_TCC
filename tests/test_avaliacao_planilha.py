"""Planilha cega da rubrica: o que mostra, o que esconde, e a importação."""

import pytest
from openpyxl import load_workbook

from avaliacao_geracao import ErroDeAvaliacao
from avaliacao_geracao.planilha import (
    LIMITE_TRECHO, LinhaRubrica, gerar_planilha, gravar_csv, ler_csv, ler_planilha, validar,
)
from avaliacao_geracao.rodada import ItemDaRodada, id_do_achado, impressao
from core.analise.base import AchadoEnriquecido, Fidelidade, Geracao
from core.rules.registry import Registro
from fabricas_analise import achado, contexto, meta, trecho_de_contexto


def _gerado(objeto="T", explicacao="A tabela automática aumenta o modelo.", ctx=None) -> AchadoEnriquecido:
    return AchadoEnriquecido(
        achado=achado(objeto=objeto), contexto=ctx or contexto(2),
        geracao=Geracao(explicacao=explicacao, recomendacao="Desative a opção.", citacoes=[1]),
        situacao="gerado", cita_ancora=False, fidelidade=Fidelidade(divergencias=["42"]), modelo="m",
        segundos=8.1, tokens_prompt=999, do_cache=True,
    )


def _rodada(*itens) -> dict:
    return {id_do_achado("P1_A", i): ItemDaRodada("P1_A", i) for i in itens}


def _registro() -> Registro:
    registro = Registro()
    registro.registrar(meta(), lambda modelo: [])
    return registro


def _celulas(caminho):
    aba = load_workbook(caminho)["Rubrica"]
    return aba, [[c.value for c in linha] for linha in aba.iter_rows()]


def test_planilha_mostra_o_texto_e_esconde_os_sinais_automaticos(tmp_path):
    descartado = AchadoEnriquecido(achado=achado(objeto="X"), contexto=contexto(1), situacao="sem_fundamentacao",
                                   motivo_descarte="json_invalido", modelo="m")
    destino = tmp_path / "r.xlsx"
    assert gerar_planilha(_rodada(_gerado(), descartado), destino, registro=_registro()) == 1
    aba, celulas = _celulas(destino)
    texto = " ".join(str(v) for linha in celulas for v in linha if v is not None)
    assert "P1_A|MOD-001|T" in texto and "A tabela automática aumenta o modelo." in texto
    assert "Desative a opção." in texto and "Disable the auto date/time option." in texto
    assert "https://learn.microsoft.com/x" in texto  # referência da regra
    for proibido in ["cita_ancora", "42", "999", "8.1", "do_cache", "P1_A|MOD-001|X"]:
        assert proibido not in texto
    assert aba.protection.sheet is True
    assert "Instruções" in load_workbook(destino).sheetnames


def test_texto_longo_cortado_e_formula_vira_texto(tmp_path):
    """Review Focus 1."""
    longo = trecho_de_contexto(1, texto="x" * (LIMITE_TRECHO + 500))
    item = _gerado(explicacao="=SUM(1)", ctx=[longo, trecho_de_contexto(2)])
    destino = tmp_path / "r.xlsx"
    gerar_planilha(_rodada(item), destino)
    aba = load_workbook(destino)["Rubrica"]
    valores = [c for linha in aba.iter_rows() for c in linha if c.value is not None]
    explicacao = next(c for c in valores if c.value == "=SUM(1)")
    assert explicacao.data_type == "s"
    assert any("[…]" in str(c.value) and "x" * LIMITE_TRECHO in str(c.value) for c in valores)
    assert not any("x" * (LIMITE_TRECHO + 1) in str(c.value) for c in valores)


def test_sem_trecho_ancora_diz_ausente(tmp_path):
    """Review Focus 3."""
    item = _gerado(ctx=[trecho_de_contexto(1, origem="busca"), trecho_de_contexto(2)])
    destino = tmp_path / "r.xlsx"
    gerar_planilha(_rodada(item), destino)
    _, celulas = _celulas(destino)
    assert any("trecho âncora ausente" in str(v) for linha in celulas for v in linha)


def test_nao_sobrescreve_sem_forcar(tmp_path):
    destino = tmp_path / "r.xlsx"
    gerar_planilha(_rodada(_gerado()), destino)
    with pytest.raises(ErroDeAvaliacao, match="já existe"):
        gerar_planilha(_rodada(_gerado()), destino)
    gerar_planilha(_rodada(_gerado()), destino, forcar=True)


def _anotar(destino, **valores):
    livro = load_workbook(destino)
    aba = livro["Rubrica"]
    cabecalho = [c.value for c in next(aba.iter_rows(min_row=1, max_row=1))]
    for nome, valor in valores.items():
        aba.cell(row=2, column=cabecalho.index(nome) + 1, value=valor)
    livro.save(destino)


def test_importacao_valida_e_ida_e_volta_do_csv(tmp_path):
    item = _gerado()
    rodada = _rodada(item)
    destino = tmp_path / "r.xlsx"
    gerar_planilha(rodada, destino)
    _anotar(destino, pertinente="sim", fiel="não", recomendacao_correta="indeterminado", justificativa="inventou 42")
    lidas = ler_planilha(destino)
    assert validar(lidas, rodada, parcial=False) == []
    linhas = [l for _, l in lidas]
    gravar_csv(linhas, tmp_path / "r.csv")
    assert ler_csv(tmp_path / "r.csv") == linhas
    assert linhas[0].impressao == impressao(item)


@pytest.mark.parametrize("valores,motivo", [
    (dict(pertinente="talvez", fiel="sim", recomendacao_correta="sim"), "fora da lista"),
    (dict(pertinente="não", fiel="sim", recomendacao_correta="sim"), "sem justificativa"),
    (dict(pertinente="sim"), "sem veredito"),
])
def test_importacao_recusa(tmp_path, valores, motivo):
    rodada = _rodada(_gerado())
    destino = tmp_path / "r.xlsx"
    gerar_planilha(rodada, destino)
    _anotar(destino, **valores)
    problemas = validar(ler_planilha(destino), rodada, parcial=False)
    assert any(motivo in p.motivo for p in problemas)


def test_parcial_aceita_linha_em_branco(tmp_path):
    rodada = _rodada(_gerado())
    destino = tmp_path / "r.xlsx"
    gerar_planilha(rodada, destino)
    assert validar(ler_planilha(destino), rodada, parcial=True) == []


def test_impressao_divergente_e_id_desconhecido_sao_recusados(tmp_path):
    """Review Focus 2: texto refeito fora do cache depois da anotação."""
    destino = tmp_path / "r.xlsx"
    gerar_planilha(_rodada(_gerado()), destino)
    _anotar(destino, pertinente="sim", fiel="sim", recomendacao_correta="sim")
    lidas = ler_planilha(destino)
    refeita = _rodada(_gerado(explicacao="Outro texto."))
    assert any("impressão digital" in p.motivo for p in validar(lidas, refeita, parcial=False))
    assert any("não existe na rodada" in p.motivo for p in validar(lidas, {}, parcial=False))


def test_id_duplicado_e_recusado():
    linha = LinhaRubrica(id="a", impressao="x", pertinente="sim", fiel="sim", recomendacao_correta="sim")
    problemas = validar([(2, linha), (3, linha)], {}, parcial=False)
    assert any("duplicado" in p.motivo for p in problemas)


def test_validar_aponta_achado_gerado_ausente_da_planilha(tmp_path):
    rodada = _rodada(_gerado("A"), _gerado("B"))
    destino = tmp_path / "r.xlsx"
    gerar_planilha(rodada, destino)
    lidas = ler_planilha(destino)[:1]  # linha apagada
    ausente = sorted(rodada)[1] if lidas[0][1].id == sorted(rodada)[0] else sorted(rodada)[0]
    problemas = validar(lidas, rodada, parcial=False)
    assert [(p.linha, p.id, p.motivo) for p in problemas if p.linha is None] == [
        (None, ausente, "achado gerado da rodada ausente da planilha")]
    assert not [p for p in validar(lidas, rodada, parcial=True) if p.linha is None]


def test_planilha_mostra_a_expressao_do_achado_que_o_modelo_recebeu(tmp_path):
    """Defeito achado no teste do glossário (15/10/2026): sem a expressão, nomes que o
    modelo tirou dela pareciam inventados."""
    item = _gerado().model_copy(update={"achado": achado(objeto="T", trecho="IF([Units]<>0, [Sales]/[Units])")})
    destino = tmp_path / "r.xlsx"
    gerar_planilha(_rodada(item), destino, registro=_registro())
    _, celulas = _celulas(destino)
    assert "trecho_do_modelo" in celulas[0]
    assert "IF([Units]<>0, [Sales]/[Units])" in celulas[1]
