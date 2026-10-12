"""A planilha cega de pertinência dos complementares e o CSV oficial.

Mesmo formato do gabarito das regras (`gabarito/planilha.py`): uma aba por regra,
protegida, com só as colunas de anotação livres e lista suspensa no veredito. Nenhuma
célula diz qual busca trouxe a seção.
"""

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel

from rag.chunking import TrechoIndexado
from rag.complementares import ErroDeComplementares, ItemPool

VEREDITOS = ("pertinente", "não pertinente", "indeterminado")
PERGUNTA = "Esta seção ajuda a explicar o problema da regra ou a corrigi-lo, na mesma linha da âncora?"
LINHA_CABECALHO = 6
COLUNAS = ("id", "pagina", "secao", "titulo", "caminho", "inicio_do_texto", "url", "veredito", "justificativa")
EDITAVEIS = ("veredito", "justificativa")
LIMITE_ANCORA = 1500
LIMITE_INICIO = 600
LARGURAS = {"id": 40, "pagina": 30, "secao": 22, "titulo": 30, "caminho": 30, "inicio_do_texto": 80, "url": 30,
            "veredito": 16, "justificativa": 50}
CAMPOS_CSV = ("id", "id_regra", "id_pagina", "ancora", "veredito", "justificativa")
INSTRUCOES = [
    "Pertinência dos trechos complementares — protocolo de anotação",
    "",
    "1. Em cada aba, a âncora da regra está no topo: é a passagem que sustenta o achado.",
    "2. Julgue cada seção pela pergunta da aba: ela ajuda a explicar ou corrigir o problema da regra, na mesma linha da âncora?",
    "3. veredito: 'pertinente', 'não pertinente' ou 'indeterminado' (use a lista suspensa).",
    "4. 'indeterminado' pede justificativa.",
    "5. A planilha não diz qual busca trouxe a seção, de propósito: o julgamento é cego.",
    "6. Para conferir o andamento: python -m rag.complementares importar --parcial",
]


class LinhaPertinencia(BaseModel):
    id: str
    id_regra: str
    id_pagina: str
    ancora: str
    veredito: str | None = None
    justificativa: str | None = None


class Problema(BaseModel):
    aba: str
    linha: int | None
    id: str
    motivo: str


def _cortar(texto: str, limite: int) -> str:
    return texto if len(texto) <= limite else texto[:limite] + " […]"


def _escrever(aba, linha: int, coluna: int, valor: str):
    celula = aba.cell(row=linha, column=coluna, value=ILLEGAL_CHARACTERS_RE.sub("", valor))
    celula.data_type = "s"  # texto, sempre: um trecho que começa com `=` viraria fórmula
    celula.alignment = Alignment(wrap_text=True, vertical="top")
    return celula


def _aba(livro: Workbook, id_regra: str, itens: list[ItemPool], ancora: TrechoIndexado) -> None:
    aba = livro.create_sheet(id_regra)
    aba["A1"] = PERGUNTA
    aba["A1"].font = Font(bold=True)
    aba["A2"] = f"Âncora: {ancora.titulo_pagina} › {ancora.caminho} ({ancora.url})"
    _escrever(aba, 3, 1, _cortar(ancora.texto, LIMITE_ANCORA))
    for coluna, nome in enumerate(COLUNAS, start=1):
        aba.cell(row=LINHA_CABECALHO, column=coluna, value=nome).font = Font(bold=True)
        aba.column_dimensions[aba.cell(row=LINHA_CABECALHO, column=coluna).column_letter].width = LARGURAS[nome]
    for linha, item in enumerate(itens, start=LINHA_CABECALHO + 1):
        t = item.trecho
        valores = (item.id, t.id_pagina, t.ancora, t.titulo_pagina, t.caminho, _cortar(t.texto, LIMITE_INICIO), t.url)
        for coluna, valor in enumerate(valores, start=1):
            _escrever(aba, linha, coluna, valor)
        for nome in EDITAVEIS:
            aba.cell(row=linha, column=COLUNAS.index(nome) + 1).protection = Protection(locked=False)
    if itens:
        letra = aba.cell(row=LINHA_CABECALHO, column=COLUNAS.index("veredito") + 1).column_letter
        validacao = DataValidation(type="list", formula1=f'"{",".join(VEREDITOS)}"', allow_blank=True)
        validacao.add(f"{letra}{LINHA_CABECALHO + 1}:{letra}{LINHA_CABECALHO + len(itens)}")
        aba.add_data_validation(validacao)
    aba.freeze_panes = aba.cell(row=LINHA_CABECALHO + 1, column=2)
    aba.protection.formatColumns = False
    aba.protection.formatRows = False
    aba.protection.sheet = True


def gerar_planilha(pool: list[ItemPool], ancoras: dict[str, TrechoIndexado], destino: Path) -> None:
    if destino.exists():
        raise ErroDeComplementares(f"{destino} já existe e pode ter vereditos — apague ou renomeie antes de gerar outra")
    livro = Workbook()
    instrucoes = livro.active
    instrucoes.title = "Instruções"
    for i, linha in enumerate(INSTRUCOES, start=1):
        instrucoes.cell(row=i, column=1, value=linha)
    for id_regra in sorted({item.id_regra for item in pool}):
        _aba(livro, id_regra, [item for item in pool if item.id_regra == id_regra], ancoras[id_regra])
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)


def _texto(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor)
    return texto if texto.strip() else None


def ler_planilha(caminho: Path) -> list[tuple[str, int, LinhaPertinencia]]:
    livro = load_workbook(caminho, read_only=True)
    lidas: list[tuple[str, int, LinhaPertinencia]] = []
    for aba in livro.worksheets[1:]:
        linhas = list(aba.iter_rows(min_row=LINHA_CABECALHO, values_only=True))
        if not linhas:
            continue
        cabecalho = list(linhas[0])
        faltando = [nome for nome in COLUNAS if nome not in cabecalho]
        if faltando:
            livro.close()
            raise ErroDeComplementares(f"aba {aba.title}: coluna(s) ausente(s) no cabeçalho: {', '.join(faltando)}")
        pos = {nome: cabecalho.index(nome) for nome in COLUNAS}
        for numero, valores in enumerate(linhas[1:], start=LINHA_CABECALHO + 1):
            if valores[pos["id"]] is None:
                continue
            lidas.append((aba.title, numero, LinhaPertinencia(
                id=str(valores[pos["id"]]),
                id_regra=aba.title,
                id_pagina=str(valores[pos["pagina"]]),
                ancora=str(valores[pos["secao"]]),
                veredito=_texto(valores[pos["veredito"]]),
                justificativa=_texto(valores[pos["justificativa"]]),
            )))
    livro.close()
    return lidas


def validar(lidas: list[tuple[str, int, LinhaPertinencia]], pool: list[ItemPool], parcial: bool) -> list[Problema]:
    """Não normaliza: um veredito fora da lista é apontado, não corrigido em silêncio."""
    problemas: list[Problema] = []
    for aba, numero, linha in lidas:
        if linha.veredito is None:
            if not parcial:
                problemas.append(Problema(aba=aba, linha=numero, id=linha.id, motivo="sem veredito"))
            continue
        if linha.veredito not in VEREDITOS:
            problemas.append(Problema(aba=aba, linha=numero, id=linha.id, motivo=f"veredito fora da lista: {linha.veredito!r}"))
        elif linha.veredito == "indeterminado" and linha.justificativa is None:
            problemas.append(Problema(aba=aba, linha=numero, id=linha.id, motivo="sem justificativa"))
    na_planilha = {linha.id: (aba, numero) for aba, numero, linha in lidas}
    atuais = {item.id: item for item in pool}
    for id_, (aba, numero) in na_planilha.items():
        if id_ not in atuais:
            problemas.append(Problema(aba=aba, linha=numero, id=id_, motivo="não está mais no pool"))
    for id_, item in atuais.items():
        if id_ not in na_planilha:
            problemas.append(Problema(aba=item.id_regra, linha=None, id=id_, motivo="seção ausente da planilha"))
    return problemas


def progresso(lidas: list[tuple[str, int, LinhaPertinencia]]) -> dict[str, tuple[int, int]]:
    contagem: dict[str, tuple[int, int]] = {}
    for aba, _, linha in lidas:
        julgadas, total = contagem.get(aba, (0, 0))
        contagem[aba] = (julgadas + (linha.veredito is not None), total + 1)
    return contagem


def gravar_csv(linhas: list[LinhaPertinencia], destino: Path) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    with destino.open("w", encoding="utf-8", newline="") as arquivo:
        escritor = csv.DictWriter(arquivo, fieldnames=CAMPOS_CSV, lineterminator="\n")
        escritor.writeheader()
        for linha in sorted(linhas, key=lambda l: l.id):
            escritor.writerow({campo: getattr(linha, campo) or "" for campo in CAMPOS_CSV})


def ler_csv(caminho: Path) -> dict[str, str]:
    with caminho.open(encoding="utf-8", newline="") as arquivo:
        return {linha["id"]: linha["veredito"] for linha in csv.DictReader(arquivo)}
